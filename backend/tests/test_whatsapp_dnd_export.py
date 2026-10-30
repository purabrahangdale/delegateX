"""
Offline tests for the Global DND / Blocklist export: every matching record (not one page), split by
status, filters honoured, one row per phone, contact names only from the selected number, and no
values invented for fields a record does not hold.
"""

import csv
import io

import openpyxl
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from app.config import database
from app.whatsapp.repository import DNDRepository

from test_whatsapp_multi_number import ADMIN, add_number, hdr


@pytest.fixture
def client():
    from app.whatsapp.routes import router
    app = FastAPI()
    app.include_router(router)
    return TestClient(app)


def add(phone, reason="Manual Block", source="Manual Entry", notes=None, created_at=None):
    doc = DNDRepository.add_dnd_number({"phone_number": phone, "reason": reason, "source": source, "notes": notes})
    if created_at:
        database.whatsapp_dnd_collection.update_one({"clean_phone": doc["clean_phone"]}, {"$set": {"created_at": created_at}})
    return doc


def inbound(sender_phone, sender, number_id=None, created_at="2026-10-01T00:00:00"):
    database.whatsapp_message_collection.insert_one({
        "direction": "inbound", "sender": sender, "sender_phone": sender_phone,
        "number_id": number_id, "created_at": created_at,
    })


def csv_rows(resp):
    assert resp.status_code == 200, resp.text
    return list(csv.DictReader(io.StringIO(resp.content.decode("utf-8-sig"))))


def test_csv_exports_all_records_beyond_page_limit(client):
    for i in range(650):  # the list endpoint caps a page at 500
        add(f"+91 90000 {i:05d}")
    rows = csv_rows(client.get("/api/whatsapp/dnd/export", params={"format": "csv"}, headers=ADMIN))
    assert len(rows) == 650
    assert len({r["WhatsApp Number"] for r in rows}) == 650


def test_status_split_and_columns(client):
    add("+919000000001", reason="User Opt-out", source="Inbox Keyword", notes="Triggered by keyword 'STOP'")
    add("+919000000002", reason="Spam Complaint")
    rows = csv_rows(client.get("/api/whatsapp/dnd/export", headers=ADMIN))
    by_phone = {r["WhatsApp Number"]: r for r in rows}
    assert by_phone["+919000000001"]["Status"] == "DND"
    assert by_phone["+919000000002"]["Status"] == "Blocklist"
    assert by_phone["+919000000001"]["Notes"] == "Triggered by keyword 'STOP'"
    assert by_phone["+919000000002"]["Contact Name"] == ""  # no name on record → blank, never made up
    assert by_phone["+919000000001"]["Record ID"]
    assert "Created By" not in rows[0]

    dnd = csv_rows(client.get("/api/whatsapp/dnd/export", params={"list_type": "dnd"}, headers=ADMIN))
    assert [r["WhatsApp Number"] for r in dnd] == ["+919000000001"]
    block = csv_rows(client.get("/api/whatsapp/dnd/export", params={"list_type": "blocklist"}, headers=ADMIN))
    assert [r["WhatsApp Number"] for r in block] == ["+919000000002"]


def test_filters_and_date_range(client):
    add("+919000000001", reason="Manual Block", notes="vip complaint", created_at="2026-09-01T10:00:00")
    add("+919000000002", reason="Manual Block", created_at="2026-10-15T10:00:00")
    add("+919000000003", reason="Invalid Number", created_at="2026-10-15T11:00:00")
    add("+919000000004", reason="User Opt-out", source="Inbox Keyword", created_at="2026-10-20T23:59:59")

    rows = csv_rows(client.get("/api/whatsapp/dnd/export", params={"reason": "Manual Block"}, headers=ADMIN))
    assert {r["WhatsApp Number"] for r in rows} == {"+919000000001", "+919000000002"}
    rows = csv_rows(client.get("/api/whatsapp/dnd/export", params={"search": "vip"}, headers=ADMIN))
    assert [r["WhatsApp Number"] for r in rows] == ["+919000000001"]
    rows = csv_rows(client.get("/api/whatsapp/dnd/export",
                               params={"from_date": "2026-10-15", "to_date": "2026-10-20"}, headers=ADMIN))
    assert [r["WhatsApp Number"] for r in rows] == ["+919000000004", "+919000000003", "+919000000002"]
    # list_type and the reason filter combine instead of overriding each other
    rows = csv_rows(client.get("/api/whatsapp/dnd/export",
                               params={"list_type": "dnd", "reason": "Manual Block"}, headers=ADMIN))
    assert rows == []


def test_unblocked_and_duplicate_numbers_excluded(client):
    add("+919000000001")
    add("+91 90000 00001", reason="Spam Complaint")  # same number → upserted, still one record
    add("+919000000002")
    DNDRepository.remove_dnd("+919000000002")
    rows = csv_rows(client.get("/api/whatsapp/dnd/export", headers=ADMIN))
    assert len(rows) == 1 and rows[0]["Reason"] == "Spam Complaint"


def test_contact_names_only_from_selected_number(client):
    a = add_number("A", "555551", "222222", "token-a", phone="+91 90000 00010")
    b = add_number("B", "555552", "222222", "token-b", phone="+91 90000 00020")
    add("919000000001", reason="User Opt-out", source="Inbox Keyword")
    add("+91 90000 00002")
    inbound("919000000001", "Old Name", number_id=str(a["_id"]), created_at="2026-09-01T00:00:00")
    inbound("919000000001", "Asha", number_id=str(a["_id"]), created_at="2026-10-01T00:00:00")
    inbound("919000000002", "Ravi", number_id=str(b["_id"]))
    inbound("919000000002", "919000000002", number_id=str(a["_id"]))  # no profile name

    rows = csv_rows(client.get("/api/whatsapp/dnd/export", headers=hdr(a)))
    names = {r["WhatsApp Number"]: r["Contact Name"] for r in rows}
    assert names == {"919000000001": "Asha", "+91 90000 00002": ""}


def test_xlsx_has_sheet_per_status(client):
    add("+919000000001", reason="User Opt-out", source="Inbox Keyword")
    add("+919000000002", reason="Manual Block", notes="=HYPERLINK(\"x\")")
    resp = client.get("/api/whatsapp/dnd/export", params={"format": "xlsx"}, headers=ADMIN)
    assert resp.status_code == 200
    assert 'filename="DND_Blocklist_' in resp.headers["content-disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["DND", "Blocklist"]
    dnd = list(wb["DND"].iter_rows(values_only=True))
    block = list(wb["Blocklist"].iter_rows(values_only=True))
    assert dnd[0][1] == "WhatsApp Number" and [r[1] for r in dnd[1:]] == ["+919000000001"]
    assert [r[1] for r in block[1:]] == ["+919000000002"]
    assert block[1][5] == "'=HYPERLINK(\"x\")"  # formulas are neutralised, phone numbers are not

    resp = client.get("/api/whatsapp/dnd/export", params={"format": "xlsx", "list_type": "blocklist"}, headers=ADMIN)
    assert openpyxl.load_workbook(io.BytesIO(resp.content)).sheetnames == ["Blocklist"]


@pytest.mark.parametrize("params", [{"format": "pdf"}, {"list_type": "spam"}, {"from_date": "15-10-2026"},
                                    {"from_date": "2026-10-20", "to_date": "2026-10-01"}])
def test_invalid_params_rejected(client, params):
    assert client.get("/api/whatsapp/dnd/export", params=params, headers=ADMIN).status_code == 400


def test_list_endpoint_unchanged(client):
    add("+919000000001", reason="User Opt-out", source="Inbox Keyword", notes="stop")
    add("+919000000002")
    data = client.get("/api/whatsapp/dnd", params={"search": "stop"}).json()
    assert data["total"] == 1 and data["summary"]["total_dnd"] == 1 and data["summary"]["total_optouts"] == 1
