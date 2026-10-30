"""
Offline tests for the interactive button engagement report: clicks arrive through the real webhook path and are
attributed by `context.id` to the exact campaign message; Meta retries and status events never count as clicks;
button names come from the template and from WhatsApp (nothing hardcoded); non-responders are counted.
"""

import csv
import io
import time

import openpyxl
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from app.config import database
from app.whatsapp.services import campaign_service, webhook_service

from test_whatsapp_multi_number import MAPPING, drain, hdr, make_template, run
from test_whatsapp_multi_number import meta, no_auto_replies  # noqa: F401  (fixtures)

PNID, WABA = "111111", "222222"
BUTTONS = ["Book a visit", "Maybe later", "Opt me out"]  # deliberately not Yes / No / Interested
PHONES = ["919000000001", "919000000002", "919000000003", "919000000004"]
URL = "/api/whatsapp/reports/button-engagement"


@pytest.fixture
def client():
    from app.whatsapp.routes import router
    app = FastAPI()
    app.include_router(router)
    return TestClient(app)


@pytest.fixture
def campaign(legacy_number, meta):
    tid = make_template(WABA, "9001", name="site_visit")
    database.whatsapp_template_collection.update_one(
        {"meta_template_id": "9001"},
        {"$set": {"response_buttons": [{"type": "custom", "text": t} for t in BUTTONS]
                  + [{"type": "visit_website", "text": "Website", "url": "https://x.test"}]}})
    camp = campaign_service.create_campaign({
        "name": "Diwali offer", "template_id": tid, "variable_mapping": MAPPING,
        "recipients": [{"name": f"C{i}", "phone": p} for i, p in enumerate(PHONES)],
    }, number=legacy_number)
    campaign_service.launch_campaign(camp["_id"], {"confirm_consent": True})
    run(drain())
    jobs = {j["phone"]: j for j in database.whatsapp_campaign_recipient_collection.find({"campaign_id": camp["_id"]})}
    return {"id": camp["_id"], "number": legacy_number, "wamid": {p: jobs[p]["wamid"] for p in PHONES}}


def _deliver(message: dict, sender: str):
    payload = {"object": "whatsapp_business_account", "entry": [{"id": WABA, "changes": [{"field": "messages", "value": {
        "messaging_product": "whatsapp", "metadata": {"phone_number_id": PNID, "display_phone_number": "x"},
        "contacts": [{"wa_id": sender, "profile": {"name": "Customer"}}],
        "messages": [{"from": sender, "timestamp": str(int(time.time())), **message}]}}]}]}
    webhook_service.ingest(payload, True)
    run(webhook_service.process_pending_events())


def click(sender: str, context_wamid: str, wamid: str, text: str, payload=None):
    _deliver({"id": wamid, "type": "button", "context": {"from": "x", "id": context_wamid},
              "button": {"text": text, "payload": payload or text}}, sender)


def status(wamid: str, value: str):
    webhook_service.ingest({"object": "whatsapp_business_account", "entry": [{"id": WABA, "changes": [{"field": "messages", "value": {
        "messaging_product": "whatsapp", "metadata": {"phone_number_id": PNID},
        "statuses": [{"id": wamid, "status": value, "timestamp": str(int(time.time())), "recipient_id": "x"}]}}]}]}, True)
    run(webhook_service.process_pending_events())


def by_title(buttons):
    return {b["title"]: b for b in buttons}


def test_button_breakdown_and_no_response(campaign, client):
    w = campaign["wamid"]
    click(PHONES[0], w[PHONES[0]], "wamid.R1", "Book a visit", payload="BOOK_VISIT")
    click(PHONES[1], w[PHONES[1]], "wamid.R2", "Opt me out")
    click(PHONES[1], w[PHONES[1]], "wamid.R2", "Opt me out")  # Meta retry of the same click
    status(w[PHONES[2]], "read")  # read is not engagement

    r = client.get(URL, headers=hdr(campaign["number"]))
    assert r.status_code == 200, r.text
    c = r.json()["campaigns"][0]
    assert (c["total_recipients"], c["messages_sent"], c["total_clicks"], c["responded"], c["no_response"]) == (4, 4, 2, 2, 2)
    b = by_title(c["buttons"])
    assert list(b) == BUTTONS  # configured quick replies only, in template order, unclicked ones included
    assert b["Book a visit"]["clicks"] == 1 and b["Book a visit"]["ids"] == ["BOOK_VISIT"] and b["Book a visit"]["percentage"] == 25.0
    assert b["Maybe later"]["clicks"] == 0 and b["Opt me out"]["clicks"] == 1
    assert r.json()["summary"]["no_response"] == 2


def test_clicks_attributed_by_context_and_unknown_titles_kept(campaign, client):
    w = campaign["wamid"]
    # Contact 1 clicks twice (two distinct taps); a WhatsApp interactive reply with an id/title not in the template.
    click(PHONES[0], w[PHONES[0]], "wamid.R1", "Book a visit")
    click(PHONES[0], w[PHONES[0]], "wamid.R2", "Maybe later")
    _deliver({"id": "wamid.R3", "type": "interactive", "context": {"from": "x", "id": w[PHONES[3]]},
              "interactive": {"type": "button_reply", "button_reply": {"id": "cb_42", "title": "Call me"}}}, PHONES[3])
    # A button reply to some other (non-campaign) message is not attributed to the campaign.
    click(PHONES[2], "wamid.SOMETHING_ELSE", "wamid.R4", "Book a visit")

    c = client.get(URL, headers=hdr(campaign["number"])).json()["campaigns"][0]
    b = by_title(c["buttons"])
    assert b["Call me"] == {**b["Call me"], "clicks": 1, "ids": ["cb_42"], "configured": False}
    assert (c["total_clicks"], c["responded"], c["no_response"]) == (3, 2, 2)


def test_filters_and_rows(campaign, client):
    w = campaign["wamid"]
    click(PHONES[0], w[PHONES[0]], "wamid.R1", "Book a visit")
    click(PHONES[1], w[PHONES[1]], "wamid.R2", "Opt me out")
    h = hdr(campaign["number"])

    rows = client.get(URL, headers=h).json()["rows"]
    assert len(rows) == 4
    row = next(r for r in rows if r["WhatsApp Number"] == PHONES[0])
    assert row["Button Name"] == "Book a visit" and row["Response Status"] == "Responded"
    assert row["Message ID"] == w[PHONES[0]] and row["Reply Message ID"] == "wamid.R1" and row["Campaign Name"] == "Diwali offer"
    assert row["Template Name"] == "site_visit" and row["Click Timestamp (UTC)"]

    only = client.get(URL, headers=h, params={"response_status": "no_response"}).json()["rows"]
    assert {r["WhatsApp Number"] for r in only} == {PHONES[2], PHONES[3]}
    assert [r["WhatsApp Number"] for r in client.get(URL, headers=h, params={"button": "opt me out"}).json()["rows"]] == [PHONES[1]]
    assert client.get(URL, headers=h, params={"template": "other"}).json()["campaigns"] == []
    assert client.get(URL, headers=h, params={"from_date": "2000-01-01", "to_date": "2000-01-02"}).json()["summary"]["messages_sent"] == 0
    assert client.get(URL, headers=h, params={"response_status": "bogus"}).status_code == 400


def test_failed_messages_are_not_counted_as_reached(campaign, client):
    status(campaign["wamid"][PHONES[3]], "failed")
    c = client.get(URL, headers=hdr(campaign["number"])).json()["campaigns"][0]
    assert (c["total_recipients"], c["messages_sent"], c["no_response"]) == (4, 3, 3)


def test_button_id_recovered_from_webhook_event_for_older_replies(campaign, client):
    w = campaign["wamid"]
    click(PHONES[0], w[PHONES[0]], "wamid.R1", "Book a visit", payload="BOOK_VISIT")
    # Replies stored before this feature had no metadata.button.
    database.whatsapp_message_collection.update_one({"wamid": "wamid.R1"}, {"$unset": {"metadata.button": ""}})
    b = by_title(client.get(URL, headers=hdr(campaign["number"])).json()["campaigns"][0]["buttons"])
    assert b["Book a visit"]["ids"] == ["BOOK_VISIT"]


def test_csv_and_excel_export(campaign, client):
    w = campaign["wamid"]
    click(PHONES[0], w[PHONES[0]], "wamid.R1", "Book a visit")
    h = hdr(campaign["number"])

    r = client.get(URL + "/export", headers=h, params={"format": "csv", "response_status": "responded"})
    assert r.status_code == 200 and "Button_Engagement_" in r.headers["content-disposition"]
    rows = list(csv.DictReader(io.StringIO(r.content.decode("utf-8-sig"))))
    assert [(x["WhatsApp Number"], x["Button Name"]) for x in rows] == [(PHONES[0], "Book a visit")]

    # all_records ignores the filters
    rows = list(csv.DictReader(io.StringIO(client.get(URL + "/export", headers=h, params={
        "format": "csv", "response_status": "responded", "all_records": "true"}).content.decode("utf-8-sig"))))
    assert len(rows) == 4

    r = client.get(URL + "/export", headers=h, params={"format": "xlsx"})
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Button Summary", "Contact Responses"]
    summary = [row for row in wb["Button Summary"].iter_rows(min_row=2, values_only=True)]
    assert [(s[5], s[7], s[8]) for s in summary] == [("Book a visit", 1, 1), ("Maybe later", 0, 0), ("Opt me out", 0, 0), ("No Response", None, 3)]
    assert wb["Contact Responses"].max_row == 5
    assert client.get(URL + "/export", headers=h, params={"format": "pdf"}).status_code == 400


def test_other_numbers_campaigns_are_not_visible(campaign, client):
    from test_whatsapp_multi_number import add_number
    other = add_number("Sales", "555555", "666666", "token-c", phone="+91 90000 00003")
    assert client.get(URL, headers=hdr(other)).json()["campaigns"] == []
