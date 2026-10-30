"""
WhatsApp Automation — Interactive button engagement (campaign analytics + CSV / Excel export).

A click is an inbound `button` (template quick reply) or `interactive` (button_reply / list_reply) message whose
`context.id` is the wamid of a campaign message, so each click is attributed to the exact campaign, recipient and
message it answers. Delivery / read statuses are never engagement. Clicks are counted once per inbound wamid
(Meta retries are already deduplicated on ingest; this guards replays too).

Button names come from the template's configured quick-reply buttons plus whatever titles WhatsApp actually
returned — nothing is hardcoded. Button IDs are the ones WhatsApp sent (the payload for template quick replies).

A recipient counts as reached once Meta accepted the message (wamid stored) and it did not fail afterwards;
reached recipients without a click are "No Response".
"""

import csv
import io
from datetime import datetime, timedelta
from typing import Dict, Iterable, Iterator, List, Optional

import openpyxl
from bson import ObjectId
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill

from app.config.database import (
    whatsapp_campaign_collection as campaigns,
    whatsapp_campaign_recipient_collection as jobs,
    whatsapp_message_collection as messages,
    whatsapp_template_collection as templates,
    whatsapp_webhook_event_collection as events,
)
from app.whatsapp.services.dnd_export_service import _safe
from app.whatsapp.services.webhook_service import extract_button

RESPONDED = "Responded"
NO_RESPONSE = "No Response"
RESPONSE_STATUSES = {"responded": RESPONDED, "no_response": NO_RESPONSE}
QUICK_REPLY_TYPES = {"custom", "quick_reply"}
CLICK_MESSAGE_TYPES = ["button", "interactive"]
IN_CHUNK = 1000

HEADERS = ["Campaign Name", "Campaign ID", "Contact Name", "WhatsApp Number", "Template Name", "Button Name",
           "Button ID", "Response Status", "Click Timestamp (UTC)", "Message ID", "Reply Message ID"]
COLUMN_WIDTHS = [28, 26, 24, 18, 26, 22, 22, 16, 22, 40, 40]
SUMMARY_HEADERS = ["Campaign Name", "Campaign ID", "Template Name", "Total Recipients", "Messages Sent",
                   "Button", "Button ID(s)", "Clicks", "Contacts", "% of Recipients Reached"]
SUMMARY_WIDTHS = [28, 26, 26, 16, 14, 24, 24, 10, 10, 22]


class EngagementError(ValueError):
    pass


def _parse_date(value: Optional[str], label: str) -> Optional[datetime]:
    if not value:
        return None
    try:
        return datetime.strptime(value, "%Y-%m-%d")
    except ValueError:
        raise EngagementError(f"{label} must be a date in YYYY-MM-DD format.")


def _norm(title: Optional[str]) -> str:
    return " ".join((title or "").split()).casefold()


def _fmt_ts(value: Optional[str]) -> str:
    return (value or "").replace("T", " ")[:19]


def _chunks(items: List[str]) -> Iterator[List[str]]:
    for i in range(0, len(items), IN_CHUNK):
        yield items[i:i + IN_CHUNK]


# ═══════════════════════════════════════════════════════════════════
# DATA
# ═══════════════════════════════════════════════════════════════════

def _template_label(camp: dict) -> str:
    t = camp.get("template") or {}
    return t.get("display_name") or t.get("name") or ""


def _load_campaigns(scope_match: dict, campaign_id: Optional[str] = None, template: Optional[str] = None) -> List[dict]:
    query: dict = {**scope_match, "status": {"$ne": "draft"}}
    if campaign_id:
        if not ObjectId.is_valid(campaign_id):
            raise EngagementError("Unknown campaign.")
        query["_id"] = ObjectId(campaign_id)
    if template:
        query["$or"] = [{"template.local_id": template}, {"template.name": template}, {"template.display_name": template}]
    return list(campaigns.find(query, {"name": 1, "template": 1, "status": 1, "created_at": 1, "launched_at": 1}).sort("created_at", -1))


def _configured_buttons(camps: List[dict]) -> Dict[str, List[str]]:
    """Quick-reply button titles configured on each campaign's template, in template order."""
    ids = {str((c.get("template") or {}).get("local_id") or "") for c in camps} - {""}
    oids = [ObjectId(i) for i in ids if ObjectId.is_valid(i)]
    if not oids:
        return {}
    out = {}
    for t in templates.find({"_id": {"$in": oids}}, {"response_buttons": 1}):
        out[str(t["_id"])] = [b.get("text") for b in t.get("response_buttons") or []
                              if (b.get("type") or "").lower() in QUICK_REPLY_TYPES and b.get("text")]
    return out


def _load_recipients(campaign_ids: List[str], start: Optional[datetime], end: Optional[datetime]) -> List[dict]:
    """Every eligible recipient job; `sent` marks the ones whose message reached WhatsApp (and did not fail)."""
    query: dict = {"campaign_id": {"$in": campaign_ids}, "state": {"$nin": ["invalid", "skipped"]}}
    if start or end:
        # The date range is the campaign message's send time, so clicks and non-responders stay comparable.
        rng = {}
        if start:
            rng["$gte"] = start.isoformat()
        if end:
            rng["$lt"] = (end + timedelta(days=1)).isoformat()
        query["accepted_at"] = rng
    out = []
    for j in jobs.find(query, {"campaign_id": 1, "phone": 1, "name": 1, "wamid": 1, "state": 1, "accepted_at": 1}):
        j["sent"] = bool(j.get("wamid")) and j.get("state") != "failed"
        out.append(j)
    return out


def _load_clicks(wamids: List[str], scope_match: dict) -> Dict[str, List[dict]]:
    """Button clicks keyed by the campaign message wamid they reply to, oldest first, one per inbound message."""
    by_context: Dict[str, List[dict]] = {}
    seen = set()
    missing_ids = []
    for chunk in _chunks(wamids):
        cursor = messages.find(
            {**scope_match, "direction": "inbound", "message_type": {"$in": CLICK_MESSAGE_TYPES},
             "metadata.context.id": {"$in": chunk}},
            {"wamid": 1, "content": 1, "created_at": 1, "metadata.context": 1, "metadata.button": 1},
        )
        for m in cursor:
            key = m.get("wamid") or str(m["_id"])
            if key in seen:
                continue
            seen.add(key)
            meta = m.get("metadata") or {}
            btn = meta.get("button") or {}
            click = {
                "reply_wamid": m.get("wamid") or "",
                "title": btn.get("title") or m.get("content") or "",
                "id": btn.get("id") or "",
                "at": m.get("created_at") or "",
            }
            if not btn and click["reply_wamid"]:
                missing_ids.append(click)
            by_context.setdefault(meta["context"]["id"], []).append(click)

    # Replies stored before button IDs were persisted: recover the ID from the raw webhook event while it is retained.
    if missing_ids:
        by_wamid = {c["reply_wamid"]: c for c in missing_ids}
        for chunk in _chunks([f"message:{w}" for w in by_wamid]):
            for ev in events.find({"dedupe_key": {"$in": chunk}}, {"payload": 1}):
                btn = extract_button(ev.get("payload") or {})
                click = by_wamid.get((ev.get("payload") or {}).get("id"))
                if btn and click:
                    click["id"] = btn["id"]
                    click["title"] = btn["title"] or click["title"]

    for clicks in by_context.values():
        clicks.sort(key=lambda c: c["at"])
    return by_context


def _collect(scope_match: dict, campaign_id=None, from_date=None, to_date=None, template=None):
    start = _parse_date(from_date, "from_date")
    end = _parse_date(to_date, "to_date")
    if start and end and start > end:
        raise EngagementError("from_date must be on or before to_date.")
    camps = _load_campaigns(scope_match, campaign_id, template)
    recipients = _load_recipients([str(c["_id"]) for c in camps], start, end) if camps else []
    clicks = _load_clicks([r["wamid"] for r in recipients if r["sent"]], scope_match) if recipients else {}
    return camps, recipients, clicks


# ═══════════════════════════════════════════════════════════════════
# ANALYTICS
# ═══════════════════════════════════════════════════════════════════

class _ButtonTally:
    def __init__(self, title: str, configured: bool):
        self.title = title
        self.configured = configured
        self.ids: List[str] = []
        self.clicks = 0
        self.contacts: set = set()

    def add(self, click: dict, contact: str):
        self.clicks += 1
        self.contacts.add(contact)
        if click["id"] and click["id"] not in self.ids:
            self.ids.append(click["id"])

    def out(self, reached: int) -> dict:
        return {"title": self.title, "ids": self.ids, "configured": self.configured, "clicks": self.clicks,
                "contacts": len(self.contacts),
                "percentage": round(len(self.contacts) / reached * 100, 1) if reached else 0.0}


def _tally(recipients: List[dict], clicks: Dict[str, List[dict]], configured: List[str]) -> dict:
    buttons: Dict[str, _ButtonTally] = {}
    for title in configured:
        buttons.setdefault(_norm(title), _ButtonTally(title, True))
    total = sent = responded = total_clicks = 0
    for r in recipients:
        total += 1
        if not r["sent"]:
            continue
        sent += 1
        mine = clicks.get(r["wamid"]) or []
        if mine:
            responded += 1
        contact = r.get("phone") or str(r["_id"])
        for c in mine:
            total_clicks += 1
            buttons.setdefault(_norm(c["title"]), _ButtonTally(c["title"], False)).add(c, contact)
    return {
        "total_recipients": total,
        "messages_sent": sent,
        "responded": responded,
        "no_response": sent - responded,
        "total_clicks": total_clicks,
        "response_rate": round(responded / sent * 100, 1) if sent else 0.0,
        "no_response_rate": round((sent - responded) / sent * 100, 1) if sent else 0.0,
        "buttons": [b.out(sent) for b in buttons.values()],
    }


def build_report(scope_match: dict, campaign_id=None, from_date=None, to_date=None, template=None,
                 button=None, response_status=None, limit: int = 200, skip: int = 0) -> dict:
    """Overall + per-campaign button analytics and one page of contact-level rows."""
    camps, recipients, clicks = _collect(scope_match, campaign_id, from_date, to_date, template)
    configured = _configured_buttons(camps)
    by_campaign: Dict[str, List[dict]] = {}
    for r in recipients:
        by_campaign.setdefault(r["campaign_id"], []).append(r)

    per_campaign = []
    all_configured: List[str] = []
    for c in camps:
        cid = str(c["_id"])
        conf = configured.get(str((c.get("template") or {}).get("local_id") or ""), [])
        all_configured += [t for t in conf if _norm(t) not in {_norm(x) for x in all_configured}]
        per_campaign.append({"campaign_id": cid, "campaign_name": c.get("name") or "", "template_name": _template_label(c),
                             "status": c.get("status"), "launched_at": c.get("launched_at") or c.get("created_at"),
                             **_tally(by_campaign.get(cid, []), clicks, conf)})

    rows = list(iter_rows(camps, recipients, clicks, button=button, response_status=response_status))
    every = _load_campaigns(scope_match) if (campaign_id or template) else camps
    return {
        "summary": {"campaigns": len(camps), **_tally(recipients, clicks, all_configured)},
        "campaigns": per_campaign,
        "rows": [dict(zip(HEADERS, row)) for row in rows[skip:skip + limit]],
        "total_rows": len(rows),
        "options": {
            "campaigns": [{"id": str(c["_id"]), "name": c.get("name") or ""} for c in every],
            "templates": sorted({_template_label(c) for c in every} - {""}),
        },
    }


# ═══════════════════════════════════════════════════════════════════
# CONTACT-LEVEL ROWS & EXPORT
# ═══════════════════════════════════════════════════════════════════

def iter_rows(camps: List[dict], recipients: List[dict], clicks: Dict[str, List[dict]],
              button: Optional[str] = None, response_status: Optional[str] = None) -> Iterator[List[str]]:
    """One row per click; one 'No Response' row per reached recipient who clicked nothing."""
    if response_status and response_status not in RESPONSE_STATUSES:
        raise EngagementError(f"response_status must be one of: {', '.join(RESPONSE_STATUSES)}.")
    want_status = RESPONSE_STATUSES.get(response_status or "")
    want_button = _norm(button) if button else None
    info = {str(c["_id"]): c for c in camps}
    for r in recipients:
        if not r["sent"]:
            continue
        c = info.get(r["campaign_id"]) or {}
        base = [c.get("name") or "", r["campaign_id"], r.get("name") or "", r.get("phone") or "", _template_label(c)]
        mine = clicks.get(r["wamid"]) or []
        if not mine:
            if want_button or want_status == RESPONDED:
                continue
            yield base + ["", "", NO_RESPONSE, "", r["wamid"], ""]
            continue
        if want_status == NO_RESPONSE:
            continue
        for click in mine:
            if want_button and _norm(click["title"]) != want_button:
                continue
            yield base + [click["title"], click["id"], RESPONDED, _fmt_ts(click["at"]), r["wamid"], click["reply_wamid"]]


def export_rows(scope_match: dict, campaign_id=None, from_date=None, to_date=None, template=None,
                button=None, response_status=None):
    """(contact rows, per-campaign summary rows) for every matching record — not one page."""
    camps, recipients, clicks = _collect(scope_match, campaign_id, from_date, to_date, template)
    rows = list(iter_rows(camps, recipients, clicks, button=button, response_status=response_status))
    configured = _configured_buttons(camps)
    by_campaign: Dict[str, List[dict]] = {}
    for r in recipients:
        by_campaign.setdefault(r["campaign_id"], []).append(r)
    summary = []
    for c in camps:
        cid = str(c["_id"])
        t = _tally(by_campaign.get(cid, []), clicks, configured.get(str((c.get("template") or {}).get("local_id") or ""), []))
        base = [c.get("name") or "", cid, _template_label(c), t["total_recipients"], t["messages_sent"]]
        for b in t["buttons"]:
            summary.append(base + [b["title"], ", ".join(b["ids"]), b["clicks"], b["contacts"], b["percentage"]])
        summary.append(base + [NO_RESPONSE, "", "", t["no_response"], t["no_response_rate"]])
    return rows, summary


def stream_csv(rows: Iterable[List[str]]) -> Iterator[str]:
    """CSV text chunks; starts with a BOM so Excel opens non-ASCII names correctly."""
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    buffer.write("﻿")
    writer.writerow(HEADERS)
    for i, row in enumerate(rows, 1):
        writer.writerow([_safe(str(v)) for v in row])
        if i % IN_CHUNK == 0:
            yield buffer.getvalue()
            buffer.seek(0)
            buffer.truncate(0)
    yield buffer.getvalue()


def build_excel(rows: List[List[str]], summary: List[list]) -> bytes:
    """'Button Summary' (per campaign, per button) and 'Contact Responses' sheets, write-only to keep memory flat."""
    wb = openpyxl.Workbook(write_only=True)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="10B981", end_color="10B981", fill_type="solid")

    for title, headers, widths, data in (("Button Summary", SUMMARY_HEADERS, SUMMARY_WIDTHS, summary),
                                         ("Contact Responses", HEADERS, COLUMN_WIDTHS, rows)):
        ws = wb.create_sheet(title=title)
        for idx, width in enumerate(widths):
            ws.column_dimensions[openpyxl.utils.get_column_letter(idx + 1)].width = width
        header = []
        for h in headers:
            cell = WriteOnlyCell(ws, value=h)
            cell.font = header_font
            cell.fill = header_fill
            header.append(cell)
        ws.append(header)
        for row in data:
            ws.append([_safe(v) if isinstance(v, str) else v for v in row])

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def file_name(fmt: str) -> str:
    return f"Button_Engagement_{datetime.utcnow().strftime('%Y-%m-%d')}.{fmt}"
