"""
WhatsApp Automation — Global DND / Blocklist export (CSV or Excel).

Rows are read from `whatsapp_dnd` in batches (never the whole collection in memory) and only fields
the records actually hold are exported. Status is derived from the stored reason: a "User Opt-out"
is DND (the contact asked not to be messaged); every other reason is a Blocklist entry.
"""

import csv
import io
import re
from datetime import datetime, timedelta
from typing import Dict, Iterable, Iterator, List, Optional

import openpyxl
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill

from app.config.database import whatsapp_dnd_collection, whatsapp_message_collection
from app.whatsapp.repository import DNDRepository

DND_REASON = "User Opt-out"
LIST_TYPES = ("all", "dnd", "blocklist")
STATUS_LABELS = {"dnd": "DND", "blocklist": "Blocklist"}

HEADERS = ["Contact Name", "WhatsApp Number", "Record ID", "Status", "Reason", "Notes", "Source", "Date Added (UTC)"]
COLUMN_WIDTHS = [26, 20, 28, 12, 18, 40, 16, 22]

BATCH_SIZE = 1000

# Spreadsheet apps execute cells starting with these; phone numbers like "+91 98765 43210" are left alone.
_FORMULA_PREFIX = re.compile(r"^[=@\t\r]|^[+-](?![\d\s()-]*$)")


class ExportError(ValueError):
    pass


def _parse_date(value: Optional[str], label: str) -> Optional[datetime]:
    if not value:
        return None
    try:
        return datetime.strptime(value, "%Y-%m-%d")
    except ValueError:
        raise ExportError(f"{label} must be a date in YYYY-MM-DD format.")


def build_query(
    list_type: str = "all",
    search: Optional[str] = None,
    reason: Optional[str] = None,
    source: Optional[str] = None,
    from_date: Optional[str] = None,
    to_date: Optional[str] = None,
) -> dict:
    """Same filter as the DND list page, plus status (DND / Blocklist) and an optional Date Added range."""
    if list_type not in LIST_TYPES:
        raise ExportError(f"list_type must be one of: {', '.join(LIST_TYPES)}.")
    start = _parse_date(from_date, "from_date")
    end = _parse_date(to_date, "to_date")
    if start and end and start > end:
        raise ExportError("from_date must be on or before to_date.")

    clauses = [DNDRepository.list_query(search=search, reason=reason, source=source)]
    if list_type == "dnd":
        clauses.append({"reason": DND_REASON})
    elif list_type == "blocklist":
        clauses.append({"reason": {"$ne": DND_REASON}})

    # created_at is stored as a UTC ISO string, so string comparison orders correctly.
    created = {}
    if start:
        created["$gte"] = start.strftime("%Y-%m-%d")
    if end:
        created["$lt"] = (end + timedelta(days=1)).strftime("%Y-%m-%d")
    if created:
        clauses.append({"created_at": created})

    return clauses[0] if len(clauses) == 1 else {"$and": clauses}


def _digits(value: Optional[str]) -> str:
    return re.sub(r"\D", "", value or "")


def _contact_names(docs: List[dict], scope_match: dict) -> Dict[str, str]:
    """Latest WhatsApp profile name per phone (digits), from inbound messages of the selected number."""
    variants = set()
    for d in docs:
        clean = d.get("clean_phone")
        if clean:
            variants.update({clean, "+" + clean})
        if d.get("phone_number"):
            variants.add(d["phone_number"])
    if not variants:
        return {}

    names: Dict[str, str] = {}
    cursor = whatsapp_message_collection.find(
        {**scope_match, "direction": "inbound", "sender_phone": {"$in": list(variants)}},
        {"_id": 0, "sender": 1, "sender_phone": 1},
    ).sort("created_at", -1)
    for m in cursor:
        digits = _digits(m.get("sender_phone"))
        name = (m.get("sender") or "").strip()
        # Without a profile name the sender field holds the phone itself — that is not a name.
        if digits and name and digits not in names and _digits(name) != digits:
            names[digits] = name
    return names


def _format_date(value: Optional[str]) -> str:
    if not value:
        return ""
    try:
        return datetime.fromisoformat(value).strftime("%Y-%m-%d %H:%M:%S")
    except ValueError:
        return value


def _row(doc: dict, names: Dict[str, str]) -> List[str]:
    reason = doc.get("reason") or ""
    return [
        names.get(doc.get("clean_phone") or "", ""),
        doc.get("phone_number") or "",
        str(doc["_id"]),
        STATUS_LABELS["dnd"] if reason == DND_REASON else STATUS_LABELS["blocklist"],
        reason,
        doc.get("notes") or "",
        doc.get("source") or "",
        _format_date(doc.get("created_at")),
    ]


def iter_rows(query: dict, scope_match: dict, seen: Optional[set] = None) -> Iterator[List[str]]:
    """Yield export rows newest first, one row per phone number (pass `seen` to dedupe across calls)."""
    seen = set() if seen is None else seen
    cursor = whatsapp_dnd_collection.find(
        query,
        {"phone_number": 1, "clean_phone": 1, "reason": 1, "source": 1, "notes": 1, "created_at": 1},
    ).sort("created_at", -1).batch_size(BATCH_SIZE)

    batch: List[dict] = []
    for doc in cursor:
        key = doc.get("clean_phone") or str(doc["_id"])
        if key in seen:
            continue
        seen.add(key)
        batch.append(doc)
        if len(batch) >= BATCH_SIZE:
            names = _contact_names(batch, scope_match)
            yield from (_row(d, names) for d in batch)
            batch = []
    if batch:
        names = _contact_names(batch, scope_match)
        yield from (_row(d, names) for d in batch)


def _safe(value: str) -> str:
    return "'" + value if value and _FORMULA_PREFIX.match(value) else value


def stream_csv(rows: Iterable[List[str]]) -> Iterator[str]:
    """CSV text chunks; starts with a BOM so Excel opens non-ASCII names correctly."""
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    buffer.write("﻿")
    writer.writerow(HEADERS)
    for i, row in enumerate(rows, 1):
        writer.writerow([_safe(v) for v in row])
        if i % BATCH_SIZE == 0:
            yield buffer.getvalue()
            buffer.seek(0)
            buffer.truncate(0)
    yield buffer.getvalue()


def build_excel(sheets: List[tuple]) -> bytes:
    """`sheets` is a list of (title, rows). Uses openpyxl's write-only mode to keep memory flat."""
    wb = openpyxl.Workbook(write_only=True)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="E11D48", end_color="E11D48", fill_type="solid")

    for title, rows in sheets:
        ws = wb.create_sheet(title=title)
        for idx, width in enumerate(COLUMN_WIDTHS):
            ws.column_dimensions[openpyxl.utils.get_column_letter(idx + 1)].width = width
        header = []
        for h in HEADERS:
            cell = WriteOnlyCell(ws, value=h)
            cell.font = header_font
            cell.fill = header_fill
            header.append(cell)
        ws.append(header)
        for row in rows:
            ws.append([_safe(v) for v in row])

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def file_name(list_type: str, fmt: str) -> str:
    base = {"all": "DND_Blocklist", "dnd": "DND_Contacts", "blocklist": "Blocklist_Contacts"}[list_type]
    return f"{base}_{datetime.utcnow().strftime('%Y-%m-%d')}.{fmt}"
